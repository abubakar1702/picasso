# Copyright (c) 2026, Akash and contributors
# License: MIT. See LICENSE

import csv
import io
import json
from typing import Any

import frappe
from frappe import _
from frappe.utils import cint, cstr

STUDIO_KEY = "picasso_studio"
FEATURE_KEYS = (
	"page_transitions",
	"progress_bar",
	"ripple",
	"reveal",
	"tilt",
	"counters",
	"condensed_header",
	"back_to_top",
	"toast_timers",
	"save_pulse",
	"palette",
	"charts",
	"cards",
	"app_icons",
	"dock_autohide",
	"gradients",
	"quicklook",
	"signin_entrance",
)
DENSITIES = ("compact", "cozy", "roomy")
TOAST_POS = (
	"top-left",
	"top-center",
	"top-right",
	"bottom-left",
	"bottom-center",
	"bottom-right",
)
DOCK_CORNERS = ("br", "bl", "tr", "tl")
MAX_PREVIEW_ROWS = 40
MAX_PREVIEW_COLS = 12


def default_studio() -> dict:
	return {
		"density": "cozy",
		"motion": "on",
		"accent": "",
		"toast_position": "bottom-right",
		"dock_corner": "br",
		"features": {key: True for key in FEATURE_KEYS},
	}


def _clean_studio(raw: Any) -> dict:
	data = raw if isinstance(raw, dict) else {}
	out = default_studio()
	if data.get("density") in DENSITIES:
		out["density"] = data["density"]
	if data.get("motion") in ("on", "off"):
		out["motion"] = data["motion"]
	accent = cstr(data.get("accent") or "").strip()
	if accent.startswith("#") and len(accent) in (4, 7):
		out["accent"] = accent
	if data.get("toast_position") in TOAST_POS:
		out["toast_position"] = data["toast_position"]
	if data.get("dock_corner") in DOCK_CORNERS:
		out["dock_corner"] = data["dock_corner"]
	features = data.get("features") or {}
	if isinstance(features, dict):
		for key in FEATURE_KEYS:
			if key in features:
				out["features"][key] = bool(features[key])
	return out


def get_studio(user: str | None = None) -> dict:
	user = user or frappe.session.user
	if not user or user == "Guest":
		return default_studio()
	raw = frappe.defaults.get_user_default(STUDIO_KEY, user=user)
	if not raw:
		return default_studio()
	try:
		parsed = json.loads(raw) if isinstance(raw, str) else raw
	except (TypeError, ValueError):
		return default_studio()
	return _clean_studio(parsed)


@frappe.whitelist()
def save_studio(studio=None):
	if frappe.session.user == "Guest":
		frappe.throw(_("Not permitted"), frappe.PermissionError)
	if isinstance(studio, str):
		studio = frappe.parse_json(studio)
	clean = _clean_studio(studio)
	frappe.defaults.set_user_default(STUDIO_KEY, frappe.as_json(clean))
	return clean


@frappe.whitelist()
def peek_doc(doctype: str, name: str):
	doctype = _resolve_doctype(cstr(doctype))
	name = cstr(name)
	if not doctype or not name:
		frappe.throw(_("Missing document"))
	if not frappe.has_permission(doctype, "read", name):
		frappe.throw(_("Not permitted"), frappe.PermissionError)

	doc = frappe.get_doc(doctype, name)
	meta = frappe.get_meta(doctype)
	title = doc.get_title() if hasattr(doc, "get_title") else name
	fields = []
	for df in meta.fields:
		if df.hidden or df.fieldtype in (
			"Section Break",
			"Column Break",
			"Tab Break",
			"Table",
			"Table MultiSelect",
			"HTML",
			"Button",
			"Fold",
			"Heading",
			"Password",
			"Attach",
			"Attach Image",
			"Signature",
			"Code",
			"Text Editor",
			"HTML Editor",
		):
			continue
		if not doc.has_permlevel_access_to(df.fieldname, df):
			continue
		value = doc.get(df.fieldname)
		if value in (None, ""):
			continue
		fields.append({"label": _(df.label or df.fieldname), "value": cstr(value)})
		if len(fields) >= 8:
			break

	return {
		"doctype": doctype,
		"name": name,
		"title": cstr(title) or name,
		"modified": cstr(doc.modified),
		"fields": fields,
		"route": f"/app/{frappe.scrub(doctype)}/{name}",
	}


def _resolve_doctype(value: str) -> str:
	if frappe.db.exists("DocType", value):
		return value
	guess = value.replace("-", " ").title()
	if frappe.db.exists("DocType", guess):
		return guess
	name = frappe.db.get_value("DocType", {"name": ["like", value.replace("-", "%")]}, "name")
	return name or value


def _get_file_doc(file_url: str):
	file_url = cstr(file_url)
	if not file_url:
		frappe.throw(_("Missing file"))
	name = frappe.db.get_value("File", {"file_url": file_url}, "name")
	if not name:
		frappe.throw(_("File not found"), frappe.DoesNotExistError)
	doc = frappe.get_doc("File", name)
	doc.check_permission("read")
	return doc


@frappe.whitelist()
def peek_file(file_url: str):
	doc = _get_file_doc(file_url)
	ext = (doc.file_name or "").rsplit(".", 1)
	ext = ext[-1].lower() if len(ext) == 2 else ""
	payload = {
		"name": doc.file_name,
		"url": doc.file_url,
		"size": cint(doc.file_size),
		"ext": ext,
		"kind": "file",
	}
	if ext in ("csv", "txt"):
		payload["kind"] = "table"
		payload["sheets"] = [{"name": doc.file_name, "rows": _read_csv(doc)}]
	elif ext in ("xlsx", "xlsm"):
		payload["kind"] = "table"
		payload["sheets"] = _read_xlsx(doc)
	elif ext in ("png", "jpg", "jpeg", "gif", "webp", "svg", "bmp"):
		payload["kind"] = "image"
	elif ext == "pdf":
		payload["kind"] = "pdf"
	elif ext in ("mp4", "webm"):
		payload["kind"] = "video"
	elif ext in ("mp3", "wav", "ogg"):
		payload["kind"] = "audio"
	return payload


def _file_bytes(doc) -> bytes:
	content = doc.get_content()
	if isinstance(content, str):
		return content.encode("utf-8", errors="replace")
	return content or b""


def _read_csv(doc) -> list[list[str]]:
	text = _file_bytes(doc).decode("utf-8", errors="replace")
	reader = csv.reader(io.StringIO(text))
	rows = []
	for i, row in enumerate(reader):
		if i >= MAX_PREVIEW_ROWS:
			break
		rows.append([cstr(cell)[:200] for cell in row[:MAX_PREVIEW_COLS]])
	return rows


def _read_xlsx(doc) -> list[dict]:
	try:
		from openpyxl import load_workbook
	except ImportError:
		return []
	bio = io.BytesIO(_file_bytes(doc))
	wb = load_workbook(bio, read_only=True, data_only=True)
	sheets = []
	for ws in wb.worksheets[:6]:
		rows = []
		for i, row in enumerate(ws.iter_rows(max_row=MAX_PREVIEW_ROWS, max_col=MAX_PREVIEW_COLS, values_only=True)):
			rows.append([cstr(cell)[:200] if cell is not None else "" for cell in row])
		sheets.append({"name": ws.title, "rows": rows})
	wb.close()
	return sheets
